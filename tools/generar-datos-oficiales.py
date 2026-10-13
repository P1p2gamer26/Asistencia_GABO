"""Convierte los archivos oficiales del colegio en lo que sabe leer el sistema.

    python tools/generar-datos-oficiales.py

Lee de docs/datos-oficiales/ y escribe alli mismo:
  - estudiantes.csv y horario.csv, con las cabeceras exactas del panel de carga
    (/admin -> Carga de datos, ImportService).
  - reset-produccion.sql: un solo archivo para el SQL editor de Supabase que deja la
    base desde cero con todo lo de arriba (limpieza + admin + materias + docentes con su
    correo real + estudiantes + horario).
Los tres llevan datos personales de menores: estan en .gitignore, no se versionan.

No toca la base.
"""
import collections
import csv
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
import pdfplumber

RAIZ = Path(__file__).resolve().parent.parent / "docs" / "datos-oficiales"
PLANO = RAIZ / "Plano 17 09 26.xlsx"
PLANTA = RAIZ / "PLANTA07092026.xlsx"        # hoja "Hoja1": docente -> correo real
PDF_CURSOS = RAIZ / "horario-cursos-2026-2.pdf"
PDF_CURSOS_PRIMARIA = RAIZ / "HORARIOS CURSOS PRIMARIA 2026 1-5.pdf"
PDF_DOCENTES_PRIMARIA = [
    RAIZ / "HORARIO DOCENTES PRIMERO OFICIAL-1.pdf",
    RAIZ / "HORARIO_DOCENTES__SEGUNDO-1.pdf",
    RAIZ / "DOCENTES FINAL PRIMARIA 3-5 Y ESPECIALIDADES (1)-2.pdf",
]
CLAVE_TEMPORAL = "$2a$10$Dj7iHjr8j08eQUlmQcVd5uM9.8ffEMX0WtxdQPz3IAsepUn6jQnTu"  # cambiar123

# ---------------------------------------------------------------- estudiantes


def nombre_curso(grado, grupo):
    """GRADO_COD/GRUPO del plano -> el `grade` que usa el sistema.

    Los grados 1..11 usan el GRUPO tal cual ('601', '1001'): es el mismo codigo que
    traen los PDF de horario. Preescolar viene con signo negativo en el plano y no se
    le puede quitar el signo sin mas ('-101' de jardin chocaria con el '101' de
    primero), asi que lleva prefijo de letra.
    """
    grado, grupo = str(grado).strip(), str(grupo).strip()
    if grado == "0":                       # transicion: GRUPO es 1, 2, 3
        return "T0" + grupo
    if grado == "-1":                      # jardin: -101, -102, -103
        return "J" + grupo[-2:]
    if grado == "-2":                      # prejardin: -201
        return "PJ" + grupo[-2:]
    return grupo                           # 1..11 y aceleracion (99xx, PB01, SA301)


def estudiantes():
    hoja = openpyxl.load_workbook(PLANO, read_only=True).active
    filas = hoja.iter_rows(values_only=True)
    col = {n: i for i, n in enumerate(next(filas))}

    def campo(fila, nombre):
        v = fila[col[nombre]]
        return "" if v is None else str(v).strip()

    def activo(fila):
        """'ESTADO' del plano -> active. 'RETIRADO' es el unico estado que saca del
        sistema: 'MATRICULADO' y 'ASIGNADO' siguen siendo estudiantes activos. Si el
        plano no trae la columna, todos quedan activos (como antes)."""
        c = campo(fila, "ESTADO")
        if not c:
            return True
        return c != "RETIRADO"

    salida, cursos, vistos = [], collections.Counter(), set()
    for n, fila in enumerate(filas, start=2):
        doc = campo(fila, "DOC")
        if not doc:
            continue
        if doc in vistos:
            sys.exit("Fila %d: documento repetido %s" % (n, doc))
        vistos.add(doc)
        curso = nombre_curso(campo(fila, "GRADO_COD"), campo(fila, "GRUPO"))
        campos = [doc, campo(fila, "NOMBRE1"), campo(fila, "NOMBRE2"),
                  campo(fila, "APELLIDO1"), campo(fila, "APELLIDO2"), curso]
        # El importador parte por comas a secas: una coma dentro de un nombre correria
        # las columnas en silencio. Mejor parar aqui.
        if any("," in c for c in campos):
            sys.exit("Fila %d: hay una coma dentro de un campo: %s" % (n, campos))
        if not campos[1] or not campos[3]:
            sys.exit("Fila %d: falta nombre o apellido: %s" % (n, campos))
        campos.append("true" if activo(fila) else "false")
        salida.append(campos)
        cursos[curso] += 1
    return salida, cursos


# ------------------------------------------------------------------- horario
# El PDF trae una pagina por curso: 5 renglones (Dia 1..5) x 8 columnas (los bloques
# 1..7 mas el DESCANSO). Las materias y los docentes se resuelven contra una lista
# cerrada, sin tildes como en la semilla: el PDF corta las palabras que no caben en la
# celda ("INFORM" + "ATICA", "CIENCIAS POLITICAS Y E") y hay que volver a pegarlas. Si
# algo no calza, el script para: mejor eso que meterle a la base una materia inventada.

MATERIAS = [
    "ARTES", "BIOLOGIA", "CALCULO", "CIENCIAS NATURALES",
    "CIENCIAS POLITICAS Y ECONOMICAS", "DANZAS", "EDU. FISICA", "ESPANOL", "ETICA",
    "FILOSOFIA", "FISICA", "INFORMATICA", "INGLES", "LECTURA CRITICA", "LENGUAJE", "MATEMATICAS",
    "MUSICA", "PROFUNDIZACION", "QUIMICA", "RELIGION", "SOCIALES", "TEATRO",
    "TECNOLOGIA", "TRIGONOMETRIA",
]

# Como aparecen en el PDF de cursos, ya sin tildes. Varios salen con un solo nombre
# ("YAMILE", "ANDERSSON"): asi los imprime el horario.
DOCENTES = [
    "ALEJANDRO SANCHEZ", "ALEXANDER BELTRAN", "ANDERSSON", "ANDERSSON TUNJANO",
    "CARMEN GUERRERO", "CAROLINA SALAMANCA", "CELIMA IBARRA", "CESAR ZARATE",
    "CONSTANZA ARREDONDO", "FABIOLA RICO", "FRANCISCO PALACIO", "FREDDY FARFAN",
    "GISELA PINILLA", "GUSTAVO MORENO", "JAIRO MUNOZ", "JOHN NOMESQUI",
    "JUAN C RINCON", "JULIAN VARGAS", "LEIDY LINARES", "LINA GIRALDO",
    "MARCELA PARRA", "MARCELA VERA", "MELQUISEDEC SARMIENTO", "MIGUEL BACCA",
    "MILLER VILLARRAGA", "OSCAR RAMIREZ", "VICTOR ESPITIA", "YAMILE",
    "YAMILE PACHON", "YAZMIN MORENO",
]

# El horario imprime a algunos docentes con un solo nombre. Sin esto, "YAMILE" y
# "YAMILE PACHON" serian dos cuentas distintas para la misma persona.
COMPLETOS = {"YAMILE": "YAMILE PACHON", "ANDERSSON": "ANDERSSON TUNJANO"}

# El PDF identifica a los docentes por nombre y el sistema por correo. Los correos
# reales vienen en la hoja "Hoja1" de la planta ("APELLIDO1 APELLIDO2 NOMBRE1 NOMBRE2",
# correo). Un nombre del PDF calza con una fila si todos sus tokens de mas de una letra
# ("JOAN C WILCHES" -> JOAN, WILCHES) estan en el nombre de la planta; tiene que calzar
# exactamente una. Lo que no calce se arregla a mano aqui.
EXCEPCIONES = {"JOHN SALAMANCA": "SALAMANCA RUBIO JHON ALVARO"}
# Correos que la planta no trae y el colegio entrego despues (WhatsApp de coordinacion).
CORREOS_ENTREGADOS = {"YAMILE PACHON": "aida.pachon414@educacionbogota.edu.co"}  # 13-oct-2026


def planta():
    hoja = openpyxl.load_workbook(PLANTA, read_only=True)["Hoja1"]
    filas = list(hoja.iter_rows(values_only=True))[1:]
    return [(plano(str(n)), (c or "").strip().lower()) for n, c in filas if n]


def correo(nombre, filas):
    """Nombre del PDF -> correo. Devuelve tambien si es provisional: sin correo en la
    planta queda inicial + apellido @ggm.edu.co, como antes de tener la planta."""
    completo = COMPLETOS.get(nombre, nombre)
    if completo in EXCEPCIONES:
        calzan = [f for f in filas if f[0] == EXCEPCIONES[completo]]
    else:
        tokens = [t for t in completo.split() if len(t) > 1]
        calzan = [f for f in filas if all(t in f[0].split() for t in tokens)]
    if len(calzan) != 1:
        sys.exit("docente %r: calzan %d filas de la planta %s"
                 % (completo, len(calzan), [f[0] for f in calzan]))
    if calzan[0][1]:
        return calzan[0][1], False
    if completo in CORREOS_ENTREGADOS:
        return CORREOS_ENTREGADOS[completo], False
    partes = completo.lower().split()
    return "%s%s@ggm.edu.co" % (partes[0][0], partes[-1]), True


BLOQUES = [                                # columna del PDF -> bloque y horas
    (1, "06:30", "07:35"), (2, "07:35", "08:25"), (3, "08:25", "09:15"),
    (4, "09:15", "10:05"), None,           # la quinta columna es el DESCANSO
    (5, "10:35", "11:25"), (6, "11:25", "12:15"), (7, "12:15", "13:30"),
]
COLUMNAS = [90, 175, 262, 348, 434, 521, 607, 693, 779]
# Primaria: 6 bloques, descanso despues del tercero, columnas mas anchas.
BLOQUES_PRIMARIA = [
    (1, "06:30", "07:35"), (2, "07:35", "08:25"), (3, "08:25", "09:15"), None,
    (4, "09:45", "10:35"), (5, "10:35", "11:25"), (6, "11:25", "12:30"),
]
COLUMNAS_PRIMARIA = [88, 187, 286, 385, 484, 582, 681, 779]


def plano(texto):
    """'INGLES' <- 'Ingles' / 'INGLES' con tilde. Los nombres van sin tildes, como en
    la semilla (V3__datos_semilla.sql: 'Matematicas', 'Espanol')."""
    sin = unicodedata.normalize("NFD", texto)
    return re.sub(r"\s+", " ", "".join(c for c in sin if not unicodedata.combining(c))
                  ).strip().upper()


def resolver(texto, candidatos, donde):
    t = plano(texto)
    iguales = [c for c in candidatos if c == t] or [c for c in candidatos
                                                    if c.startswith(t)]
    if len(iguales) != 1:
        sys.exit("%s: no se pudo resolver %r (calzan %s)" % (donde, texto, iguales))
    return iguales[0]


def unir(lineas, candidatos, donde):
    """Une los renglones de una celda. El PDF parte las palabras largas por la mitad
    ('INFORM' + 'ATICA'), asi que se prueba pegado y separado contra la lista."""
    for pegamento in ("", " "):
        texto = pegamento.join(lineas)
        if any(c.startswith(plano(texto)) for c in candidatos):
            return resolver(texto, candidatos, donde)
    return resolver(" ".join(lineas), candidatos, donde)


def bandas(pagina):
    """Los cinco renglones Dia 1..5, por las lineas horizontales de la tabla."""
    ys = sorted({round(r["top"], 1) for r in pagina.rects
                 if r["width"] > 60 and r["height"] < 2})
    return [(ys[i], ys[i + 1]) for i in range(1, len(ys) - 1)]


def separadores(pagina, arriba, abajo, columnas=COLUMNAS):
    """Las verticales presentes en ese renglon. Las que faltan son celdas unidas: una
    clase que ocupa dos bloques seguidos. Se devuelven ajustadas a la columna mas
    cercana: los PDF dibujan la linea 1 o 2 puntos corridos segun el archivo."""
    alto = collections.Counter()
    for r in pagina.rects:
        if (r["width"] < 3 and r["height"] > 3
                and r["top"] >= arriba - 1 and r["bottom"] <= abajo + 1):
            alto[round(r["x0"])] += r["height"]
    presentes = [x for x, h in alto.items() if h > (abajo - arriba) * 0.8 and x >= 80]
    return sorted({min(columnas, key=lambda c: abs(c - x)) for x in presentes}
                  | {columnas[-1]})


def renglones(palabras):
    grupos = collections.OrderedDict()
    for p in sorted(palabras, key=lambda p: (round(p["top"] / 4), p["x0"])):
        grupos.setdefault(round(p["top"] / 4), []).append(p["text"])
    return [" ".join(v) for v in grupos.values()]


def primer_docente(crudo, donde):
    """La celda de profundizacion de 10 y 11 lista dos electivas con su docente cada
    una, todo en letra pequena. El sistema guarda un docente por bloque: queda el
    primero que aparezca."""
    texto = plano(crudo)
    encontrados = [(texto.index(d), -len(d), d) for d in DOCENTES if d in texto]
    if not encontrados:
        sys.exit("%s: profundizacion sin docente reconocible: %r" % (donde, texto))
    return min(encontrados)[2]


def titulo(nombre):
    """MATEMATICAS -> Matematicas, como los nombres de materia de la semilla."""
    return " ".join(p[:1] + p[1:].lower() for p in nombre.split())


def horario(ruta=PDF_CURSOS, columnas=COLUMNAS, bloques=BLOQUES, docentes=DOCENTES):
    """PDF por curso -> [(curso, dia, bloque, inicio, fin, materia, docente, aula)]."""
    salida = []
    with pdfplumber.open(ruta) as pdf:
        for pagina in pdf.pages:
            curso = pagina.extract_text().split("\n")[0].strip()
            for dia, (arriba, abajo) in enumerate(bandas(pagina), start=1):
                xs = separadores(pagina, arriba, abajo, columnas)
                for a, b in zip(xs, xs[1:]):
                    ia, ib = columnas.index(a), columnas.index(b)
                    if ib == ia + 1 and bloques[ia] is None:
                        continue                        # DESCANSO
                    palabras = pagina.crop(
                        (a + 1, arriba + 1, b - 1, abajo - 1)
                    ).extract_words(extra_attrs=["size"])
                    if not palabras:
                        continue                        # bloque libre
                    # Dos celdas seguidas con el mismo texto son dos bloques de la misma
                    # clase (803 dia 4, profundizacion en el bloque 7), no un duplicado:
                    # saltarlas se comia 29 bloques. comprobar() vigila las ranuras repetidas.
                    donde = "%s dia %d x=%d" % (curso, dia, a)
                    # La materia es lo mas grande de la celda (12pt en primaria, donde el
                    # docente va a 10.6 y no cabe un umbral fijo).
                    mayor = max(p["size"] for p in palabras)
                    grandes = [p for p in palabras if p["size"] > 10
                               and p["size"] >= mayor - 0.5]
                    if grandes:
                        lineas = renglones(grandes)
                        if lineas[0].strip().upper() == "OPTATIVA":
                            lineas = lineas[1:]         # el rotulo, no la materia
                        materia = unir(lineas, MATERIAS, donde)
                        y0 = min(p["top"] for p in grandes)
                        y1 = max(p["bottom"] for p in grandes)
                        docente = unir(renglones([p for p in palabras
                                                  if p["size"] < mayor - 0.5 and p["bottom"] <= y0]),
                                       docentes, donde)
                        aula = " ".join(renglones([p for p in palabras
                                                   if p["size"] < mayor - 0.5 and p["top"] >= y1]))
                    else:
                        materia, aula = "PROFUNDIZACION", ""
                        docente = primer_docente(" ".join(renglones(palabras)), donde)
                    for col in range(ia, ib):
                        if bloques[col] is None:
                            continue
                        n, ini, fin = bloques[col]
                        salida.append((curso, dia, n, ini, fin, titulo(materia), docente,
                                       plano(aula).replace(",", " ")))
    return salida


# Los PDF por docente de primaria: una pagina por docente ("Profesor NOMBRE"), misma
# tabla Dia 1..5 x 7 columnas, pero cada celda trae el CODIGO DEL CURSO en letra grande
# (3 digitos, >= 12pt) y alrededor, en letra chica, la materia y el aula (arriba o
# abajo, segun el archivo). Las celdas que no son clase (atencion a familias,
# acompanamiento de descanso, reunion de area) no traen numero y se ignoran.
AULA = re.compile(r"LAB\.? BIOLOGIA PRIA|EDU\.? FISICA \d|ARTES PRIM \d?|TECNOLOGIA PRIM"
                  r"|IFORMATICA \d?|\b\d{1,2}\b")


def horario_docentes_primaria():
    salida, nombres = [], []
    columnas, bloques = COLUMNAS_PRIMARIA, BLOQUES_PRIMARIA
    for ruta in PDF_DOCENTES_PRIMARIA:
        with pdfplumber.open(ruta) as pdf:
            for pagina in pdf.pages:
                primera = pagina.extract_text().split("\n")[0]
                if not primera.startswith("Profesor"):
                    continue                            # pagina vacia al final
                docente = plano(primera.split(" ", 1)[1])
                nombres.append(docente)
                for dia, (arriba, abajo) in enumerate(bandas(pagina), start=1):
                    palabras = pagina.crop((columnas[0], arriba, columnas[-1], abajo)
                                           ).extract_words(extra_attrs=["size"])
                    celdas = []                         # (curso, ia, ib, materias, texto)
                    for num in palabras:
                        if not (re.fullmatch(r"\d{3}", num["text"]) and num["size"] >= 12):
                            continue
                        centro = (num["x0"] + num["x1"]) / 2
                        ia = max(i for i, x in enumerate(columnas[:-1]) if x <= centro)
                        ib = ia + 1
                        for i, x in enumerate(columnas[1:-1], start=1):
                            if abs(centro - x) < 20:    # centrado sobre una linea: doble
                                ia, ib = i - 1, i + 1
                        # extract_text y no las palabras: las tildes salen como glifo
                        # aparte ("E" "TICA") y solo el texto las vuelve a pegar. El aula
                        # va antes o despues de la materia segun el archivo, y tambien
                        # nombra materias ("EDU FISICA 1", "LAB. BIOLOGIA PRIA"): se saca
                        # primero por su forma y la materia es lo mas largo que queda.
                        texto = plano(pagina.crop((columnas[ia], arriba, columnas[ib], abajo)
                                                  ).extract_text() or "")
                        # El pie "Horario generado:11/02/2026" cae dentro del ultimo renglon.
                        texto = re.sub(r"HORARIO GENERADO:\S*|ASC HORARIOS|\d+/\d+/\d+|\b\d{3}\b",
                                       " ", texto)
                        aula = " ".join(AULA.findall(texto))
                        resto = AULA.sub(" ", texto)
                        materias = sorted((m for m in MATERIAS
                                           if re.search(r"\b%s\b" % re.escape(m), resto)),
                                          key=len, reverse=True)
                        celdas.append((num["text"], ia, ib, materias, aula))
                    for curso, ia, ib, materias, aula in celdas:
                        if not materias:
                            # Clase doble que cruza el descanso: aSc imprime el curso en
                            # las dos celdas y la materia en una sola. Se toma la de la
                            # celda vecina del mismo curso en el mismo renglon.
                            vecinas = [c for c in celdas if c[0] == curso and c[3]]
                            if not vecinas:
                                sys.exit("%s dia %d curso %s: sin materia reconocible"
                                         % (docente, dia, curso))
                            materias, aula = vecinas[0][3], vecinas[0][4]
                        materia = materias[0]
                        for col in range(ia, ib):
                            if bloques[col] is None:
                                continue
                            n, ini, fin = bloques[col]
                            salida.append((curso, dia, n, ini, fin, titulo(materia),
                                           docente, aula))
    return salida, nombres


def combinar(fuentes):
    """Une varias listas de bloques; ante el mismo slot gana la fuente que va primero.
    Los PDF de cursos y de docentes son de fechas distintas y no siempre coinciden: se
    avisa de cada choque para cotejarlo con el colegio, pero no se para."""
    por_slot, choques = {}, []
    for bloques in fuentes:
        for b in bloques:
            otro = por_slot.setdefault(b[:3], b)
            if otro is not b and otro[5:7] != b[5:7]:
                choques.append("%s dia %d bloque %d: queda %s (%s); el otro PDF dice %s (%s)"
                               % (b[0], b[1], b[2], otro[5], otro[6], b[5], b[6]))
    if choques:
        print("  OJO: los PDF no coinciden en %d bloques (gana el mas reciente):\n    %s"
              % (len(choques), "\n    ".join(choques)))
    return list(por_slot.values())


# ------------------------------------------------------------------- salida


def es_primaria(curso):
    return curso.isdigit() and len(curso) == 3 and curso < "6"   # 101..504


def comprobar(bloques, cursos):
    """Lo que tiene que cumplirse si los PDF se leyeron bien. Si algo no cuadra es mejor
    saberlo aqui que despues de subirle a produccion un horario mal armado."""
    ranuras = collections.Counter((b[0], b[1], b[2]) for b in bloques)
    repetidas = [r for r, n in ranuras.items() if n > 1]
    assert not repetidas, "ranuras repetidas (se pisarian al importar): %s" % repetidas[:5]
    for (curso, dia), n in collections.Counter((b[0], b[1]) for b in bloques).items():
        tope = 6 if es_primaria(curso) else 7
        assert n <= tope, "%s dia %d quedo con %d bloques (tope %d)" % (curso, dia, n, tope)
    sin_curso = {b[0] for b in bloques} - set(cursos)
    assert not sin_curso, "hay horario de cursos que no existen en el plano: %s" % sin_curso
    ocupado = collections.defaultdict(set)
    for b in bloques:
        ocupado[(b[6], b[1], b[2])].add(b[0])
    dobles = sorted((k, sorted(v)) for k, v in ocupado.items() if len(v) > 1)
    if dobles:
        print("  OJO: docentes en dos cursos a la vez (cotejar con el colegio): %s"
              % "; ".join("%s dia %d bloque %d %s" % (k + (v,)) for k, v in dobles))
    # Primaria deberia tener los 30 (6 x 5) llenos. Un hueco puede ser real (el PDF lo
    # trae vacio) o una celda que no se leyo: se avisa, no se para.
    for curso in sorted(c for c in cursos if es_primaria(c)):
        llenos = {(b[1], b[2]) for b in bloques if b[0] == curso}
        huecos = [(d, n) for d in range(1, 6) for n in range(1, 7) if (d, n) not in llenos]
        if huecos:
            print("  OJO %s: %d huecos %s" % (curso, len(huecos), huecos))


# --------------------------------------------------------------------- salida


def escribir(ruta, cabecera, filas):
    with open(ruta, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(cabecera)
        w.writerows(filas)
    print("%s: %d lineas" % (ruta.name, len(filas)))


def q(v):
    """Literal SQL entre comillas simples."""
    return "'" + str(v).replace("'", "''") + "'"


def escribir_sql(ruta, alumnos, bloques, docentes):
    """Un solo archivo para el SQL editor de Supabase: la misma guarda y limpieza de
    tools/limpiar-produccion.sql y despues todo lo nuevo. Se corre entero o no se corre."""
    materias = sorted({b[5] for b in bloques})
    borrar = ("notifications", "attendance", "entry_log", "guardianships",
              "schedule_blocks", "students", "subjects")
    L = ["-- Base de produccion desde cero con los datos oficiales del colegio:",
         "-- %d estudiantes, %d docentes, %d materias, %d bloques."
         % (len(alumnos), len(docentes), len(materias), len(bloques)),
         "-- Generado por tools/generar-datos-oficiales.py, no editar a mano. Lleva datos",
         "-- de menores: no se versiona.", "--",
         "-- Descomenta la linea de abajo solo cuando la conexion apunte a produccion.",
         "-- SET app.borrar_todo = 'si';",
         "DO $$ BEGIN",
         "  IF coalesce(current_setting('app.borrar_todo', true), '') <> 'si' THEN",
         "    RAISE EXCEPTION 'Detenido: esto BORRA estudiantes, horario y asistencia de "
         "la base %. Descomenta el SET app.borrar_todo y vuelve a correrlo.', "
         "current_database();",
         "  END IF;", "END $$;", "", "BEGIN;", "",
         "-- 1. Limpieza (mismo orden que tools/limpiar-produccion.sql)"]
    L += ["DELETE FROM %s;" % t for t in borrar]
    L += ["DELETE FROM users WHERE role <> 'ADMIN';", "",
          "-- 2. Coordinacion",
          "INSERT INTO users (email, password_hash, full_name, role, must_change_password)",
          "VALUES ('mequinterog@educacionbogota.edu.co', %s, 'Coordinacion GGM', 'ADMIN', TRUE)"
          % q(CLAVE_TEMPORAL),
          "ON CONFLICT (email) DO NOTHING;", "",
          "-- 3. Materias",
          "INSERT INTO subjects (name) VALUES",
          ",\n".join("  (%s)" % q(m) for m in materias),
          "ON CONFLICT (name) DO NOTHING;", "",
          "-- 4. Docentes (clave temporal cambiar123; la aplicacion obliga a cambiarla)",
          "INSERT INTO users (email, password_hash, full_name, role, must_change_password) VALUES",
          ",\n".join("  (%s, %s, %s, 'DOCENTE', TRUE)" % (q(c), q(CLAVE_TEMPORAL), q(n))
                     for n, c in sorted(docentes.items())),
          "ON CONFLICT (email) DO UPDATE SET full_name = EXCLUDED.full_name;", "",
          "-- 5. Estudiantes",
          "INSERT INTO students (document_id, first_name, middle_name, last_name, "
          "second_surname, grade, active) VALUES",
          ",\n".join("  (%s, %s, %s, %s, %s, %s, %s)"
                     % (q(a[0]), q(a[1]), q(a[2]) if a[2] else "NULL", q(a[3]),
                        q(a[4]) if a[4] else "NULL", q(a[5]), a[6].upper())
                     for a in alumnos),
          "ON CONFLICT (document_id) DO UPDATE SET first_name = EXCLUDED.first_name,",
          "  middle_name = EXCLUDED.middle_name, last_name = EXCLUDED.last_name,",
          "  second_surname = EXCLUDED.second_surname, grade = EXCLUDED.grade,",
          "  active = EXCLUDED.active;", "",
          "-- 6. Horario",
          "INSERT INTO schedule_blocks (grade, weekday, block_no, start_time, end_time, "
          "subject_id, teacher_id, room)",
          "SELECT b.grade, b.weekday, b.block_no, b.start_time, b.end_time, s.id, u.id, "
          "nullif(b.room, '')",
          "FROM (VALUES",
          ",\n".join("  (%s, %d, %d, '%s'::time, '%s'::time, %s, %s, %s)"
                     % (q(b[0]), b[1], b[2], b[3], b[4], q(b[5]), q(b[6]), q(b[7]))
                     for b in bloques),
          ") AS b(grade, weekday, block_no, start_time, end_time, subject, teacher_email, room)",
          "JOIN subjects s ON s.name = b.subject",
          "JOIN users u ON u.email = b.teacher_email",
          "ON CONFLICT (grade, weekday, block_no) DO UPDATE",
          "  SET start_time = EXCLUDED.start_time, end_time = EXCLUDED.end_time,",
          "      subject_id = EXCLUDED.subject_id, teacher_id = EXCLUDED.teacher_id,",
          "      room = EXCLUDED.room;", "", "COMMIT;", "",
          "-- Verificacion: tiene que dar %d, %d, %d, %d"
          % (len(alumnos), len(docentes), len(materias), len(bloques)),
          "SELECT (SELECT count(*) FROM students) AS estudiantes,",
          "       (SELECT count(*) FROM users WHERE role = 'DOCENTE') AS docentes,",
          "       (SELECT count(*) FROM subjects) AS materias,",
          "       (SELECT count(*) FROM schedule_blocks) AS bloques;", ""]
    ruta.write_text("\n".join(L), encoding="utf-8")
    print("%s: %d lineas" % (ruta.name, len(L)))


def main():
    alumnos, cursos = estudiantes()
    escribir(RAIZ / "estudiantes.csv",
             ["document_id", "first_name", "middle_name", "last_name",
              "second_surname", "grade", "active"], alumnos)
    print("  %d cursos: %s" % (len(cursos),
                               ", ".join("%s=%d" % kv for kv in sorted(cursos.items()))))

    secundaria = horario()
    por_docente, nombres = horario_docentes_primaria()
    por_curso = horario(PDF_CURSOS_PRIMARIA, COLUMNAS_PRIMARIA, BLOQUES_PRIMARIA,
                        sorted(set(nombres)))
    # Los PDF llevan fecha: el de docentes de primero es del 14/02 y le gana al de
    # cursos (9/02); el de segundo es del 8/02 y pierde. Tercero a quinto no chocan.
    primero = [b for b in por_docente if b[0].startswith("1")]
    primaria = combinar([primero, por_curso, por_docente])
    bloques = sorted(secundaria + primaria, key=lambda b: (b[0], b[1], b[2]))
    comprobar(bloques, cursos)

    filas = planta()
    docentes, provisionales, raros = {}, {}, set()
    for b in bloques:
        nombre = COMPLETOS.get(b[6], b[6])
        if nombre in docentes:
            continue
        buzon, provisional = correo(nombre, filas)
        docentes[nombre] = buzon
        if provisional:
            provisionales[nombre] = buzon
        elif not buzon.endswith("@educacionbogota.edu.co"):
            raros.add("%s <%s>" % (nombre, buzon))
    con_correo = [b[:6] + (docentes[COMPLETOS.get(b[6], b[6])], b[7]) for b in bloques]
    escribir(RAIZ / "horario.csv",
             ["grade", "weekday", "block_no", "start_time", "end_time", "subject",
              "teacher_email", "room"], con_correo)
    escribir_sql(RAIZ / "reset-produccion.sql", alumnos, con_correo,
                 {titulo(n): c for n, c in docentes.items()})

    cuenta = collections.Counter(b[0] for b in bloques)
    print("  bloques por curso: %s"
          % ", ".join("%s=%d" % kv for kv in sorted(cuenta.items())))
    print("  resumen: %d estudiantes, %d cursos, %d docentes (%d con correo real, %d "
          "provisionales), %d bloques secundaria, %d primaria"
          % (len(alumnos), len(cursos), len(docentes), len(docentes) - len(provisionales),
             len(provisionales), len(secundaria), len(primaria)))
    if provisionales:
        print("  OJO: sin correo en la planta, quedan provisionales: %s"
              % ", ".join("%s <%s>" % kv for kv in sorted(provisionales.items())))
    if raros:
        print("  OJO: correos que no son @educacionbogota.edu.co (revisar con el "
              "colegio, aqui no se corrigen): %s" % ", ".join(sorted(raros)))


if __name__ == "__main__":
    main()
