"""Plantillas de carga para coordinación: python scripts/videos/plantillas.py
Escribe scripts/videos/plantillas/{estudiantes,horario,acudientes}.{csv,xlsx} con las cabeceras
exactas del panel de carga (/admin → Carga de datos) y filas de ejemplo inventadas.
"""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = Path(__file__).parent / "plantillas"; OUT.mkdir(exist_ok=True)

ESTUDIANTES = [
    ["document_id", "first_name", "middle_name", "last_name", "second_surname", "grade", "active"],
    ["1023456789", "Valentina", "Sofia", "Rojas", "Pineda", "601", "true"],
    ["1023456790", "Samuel", "", "Cardenas", "Mora", "601", "true"],
    ["1023456791", "Mariana", "Isabel", "Guerrero", "Lopez", "601", "true"],
    ["1023456792", "Juan", "Esteban", "Torres", "Nino", "602", "true"],
    ["1023456793", "Luciana", "", "Bermudez", "Salas", "602", "false"],
    ["1023456794", "Nicolas", "David", "Acosta", "Rueda", "701", "true"],
]
HORARIO = [["grade", "weekday", "block_no", "start_time", "end_time", "subject", "teacher_email", "room"]] + [
    ["601", d, b, ini, fin, mat, doc, aula]
    for d in ("1", "2")
    for b, ini, fin, mat, doc, aula in (
        ("1", "06:30", "07:30", "Matematicas", "docente01@ggm.edu.co", "Salon 12"),
        ("2", "07:30", "08:30", "Espanol", "docente02@ggm.edu.co", "Salon 12"),
        ("3", "08:30", "09:30", "Ciencias", "docente03@ggm.edu.co", "Laboratorio"),
        ("4", "09:45", "10:45", "Ingles", "docente04@ggm.edu.co", ""),
    )
]
ACUDIENTES = [
    ["document_id", "guardian_name", "guardian_email", "relationship"],
    ["1023456789", "Claudia Pineda", "claudia.pineda@correo.com", "Madre"],
    ["1023456790", "Jorge Cardenas", "jorge.cardenas@correo.com", "Padre"],
    ["1023456791", "Rosa Lopez", "rosa.lopez@correo.com", "Abuela"],
    ["1023456792", "Andrea Nino", "andrea.nino@correo.com", "Madre"],
]

def escribir(nombre, filas):
    with open(OUT / f"{nombre}.csv", "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(filas)
    wb = Workbook(); ws = wb.active; ws.title = nombre
    for fila in filas: ws.append(fila)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E8F0EA")
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = max(12, max(len(str(c.value or "")) for c in col) + 2)
    wb.save(OUT / f"{nombre}.xlsx")
    print(OUT / f"{nombre}.csv", OUT / f"{nombre}.xlsx")

if __name__ == "__main__":
    escribir("estudiantes", ESTUDIANTES); escribir("horario", HORARIO); escribir("acudientes", ACUDIENTES)
